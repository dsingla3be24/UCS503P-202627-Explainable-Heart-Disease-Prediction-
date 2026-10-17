"""
Generates a professional 17-slide PPTX presentation for:
  Explainable Heart Disease Prediction Using CNN-LSTM-Attention
  — Project Proposal (idea presentation, NOT completed project)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# ───────────────────── colour palette ─────────────────────
BG_DARK      = RGBColor(0x0A, 0x0E, 0x1A)
BG_MID       = RGBColor(0x0F, 0x1B, 0x2E)
BLUE_PRIMARY = RGBColor(0x37, 0x99, 0xFF)
BLUE_LIGHT   = RGBColor(0x93, 0xC5, 0xFD)
BLUE_ACCENT  = RGBColor(0x25, 0x63, 0xEB)
TEAL         = RGBColor(0x14, 0xB8, 0xA6)
TEAL_LIGHT   = RGBColor(0x5E, 0xEA, 0xD4)
RED_ACCENT   = RGBColor(0xF8, 0x71, 0x71)
RED_LIGHT    = RGBColor(0xFC, 0xA5, 0xA5)
AMBER        = RGBColor(0xFB, 0xBF, 0x24)
PURPLE       = RGBColor(0x8B, 0x5C, 0xF6)
PURPLE_LIGHT = RGBColor(0xC4, 0xB5, 0xFD)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY   = RGBColor(0xC0, 0xCF, 0xE0)
MUTED        = RGBColor(0x6B, 0x7F, 0xA3)
CARD_BORDER  = RGBColor(0x1E, 0x30, 0x4A)
TEXT_BODY    = RGBColor(0xC0, 0xCF, 0xE0)
HEADING_COL  = RGBColor(0x7E, 0xB8, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width  = SLIDE_W
prs.slide_height = SLIDE_H

# ───────────────────── helpers ─────────────────────
def set_slide_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)

def set_font(run, size=14, color=TEXT_BODY, bold=False, name='Calibri'):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = name

def add_paragraph(tf, text, size=14, color=TEXT_BODY, bold=False, align=PP_ALIGN.LEFT, space_after=Pt(6), space_before=Pt(0), name='Calibri'):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_after = space_after
    p.space_before = space_before
    run = p.add_run()
    run.text = text
    set_font(run, size, color, bold, name)
    return p

def first_paragraph(tf, text, size=14, color=TEXT_BODY, bold=False, align=PP_ALIGN.LEFT, name='Calibri'):
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(6)
    run = p.add_run()
    run.text = text
    set_font(run, size, color, bold, name)
    return p

def add_rounded_rect(slide, left, top, width, height, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_slide_number(slide, num, total=17):
    tb = add_textbox(slide, SLIDE_W - Inches(1.5), SLIDE_H - Inches(0.5), Inches(1.3), Inches(0.4))
    first_paragraph(tb.text_frame, f"{num} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)

def add_slide_header(slide, num, title):
    # Number badge
    badge = add_rounded_rect(slide, Inches(0.6), Inches(0.45), Inches(0.45), Inches(0.45),
                             BLUE_ACCENT, BLUE_ACCENT)
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = badge.text_frame.paragraphs[0].add_run()
    r.text = f"{num:02d}"
    set_font(r, 13, WHITE, True)
    badge.text_frame.word_wrap = False

    # Title
    tb = add_textbox(slide, Inches(1.2), Inches(0.35), Inches(10), Inches(0.6))
    first_paragraph(tb.text_frame, title, size=24, color=WHITE, bold=True, name='Calibri')

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.95), Inches(11.5), Pt(2))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    line.line.fill.background()

def add_proposed_tag(slide, left, top):
    tag = add_rounded_rect(slide, left, top, Inches(2.7), Inches(0.32),
                           RGBColor(0x2A, 0x1F, 0x05), AMBER)
    tag.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tag.text_frame.paragraphs[0].add_run()
    r.text = "PROPOSED — NOT YET IMPLEMENTED"
    set_font(r, 9, AMBER, True)

def add_bullet_list(tf, items, size=13, color=TEXT_BODY, bullet_color=BLUE_LIGHT):
    for item in items:
        p = tf.add_paragraph()
        p.space_after = Pt(5)
        p.level = 0
        run = p.add_run()
        run.text = f"•  {item}"
        set_font(run, size, color, name='Calibri')

def add_card(slide, left, top, width, height, title, body_lines, fill=RGBColor(0x10,0x18,0x28),
             border=RGBColor(0x1E,0x3A,0x5F), title_color=HEADING_COL, body_color=TEXT_BODY,
             title_size=13, body_size=11):
    card = add_rounded_rect(slide, left, top, width, height, fill, border)
    card.text_frame.word_wrap = True
    first_paragraph(card.text_frame, title, size=title_size, color=title_color, bold=True)
    for line in body_lines:
        add_paragraph(card.text_frame, line, size=body_size, color=body_color)
    return card

def add_flow_arrow(slide, left, top, direction='down'):
    tb = add_textbox(slide, left, top, Inches(0.4), Inches(0.3))
    first_paragraph(tb.text_frame, "↓" if direction == 'down' else "→",
                    size=16, color=BLUE_ACCENT, bold=True, align=PP_ALIGN.CENTER)

def add_flow_box(slide, left, top, width, text, fill=RGBColor(0x10,0x1C,0x30),
                 border=RGBColor(0x25,0x63,0xEB), text_color=BLUE_LIGHT, size=11):
    box = add_rounded_rect(slide, left, top, width, Inches(0.38), fill, border)
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    set_font(r, size, text_color, True)
    return box


# ════════════════════════════════════════════════════════════
#                    SLIDE 1 — TITLE
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(s, BG_DARK)

# Decorative top line
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(1.8), Inches(2.9), Pt(3))
line.fill.solid(); line.fill.fore_color.rgb = BLUE_ACCENT; line.line.fill.background()

# Title
tb = add_textbox(s, Inches(1.5), Inches(2.1), Inches(10.3), Inches(1.4))
tb.text_frame.word_wrap = True
first_paragraph(tb.text_frame, "Explainable Heart Disease Prediction", size=36, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, name='Calibri')
add_paragraph(tb.text_frame, "Using CNN-LSTM-Attention", size=36, color=BLUE_LIGHT, bold=True,
              align=PP_ALIGN.CENTER, name='Calibri')

# Subtitle
tb2 = add_textbox(s, Inches(2.5), Inches(3.6), Inches(8.3), Inches(0.6))
add_paragraph(tb2.text_frame, "A Multi-Technique XAI Framework using Attention, Grad-CAM, SHAP and LIME",
              size=16, color=HEADING_COL, align=PP_ALIGN.CENTER)

# ECG line decoration
ecg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(4.4), Inches(4.3), Pt(2))
ecg.fill.solid(); ecg.fill.fore_color.rgb = RGBColor(0x1E,0x3A,0x5F); ecg.line.fill.background()

# Meta info
tb3 = add_textbox(s, Inches(2.5), Inches(4.9), Inches(8.3), Inches(1.8))
tb3.text_frame.word_wrap = True
for line_text in [
    "Student Name  |  Roll Number  |  Section",
    "Subject: Software Engineering",
    "Faculty / Guide: Professor Name",
    "University Name",
    "August 2026"
]:
    add_paragraph(tb3.text_frame, line_text, size=13, color=MUTED, align=PP_ALIGN.CENTER,
                  space_after=Pt(4))

# PROJECT PROPOSAL badge
tag = add_rounded_rect(s, Inches(4.8), Inches(1.2), Inches(3.6), Inches(0.38),
                       RGBColor(0x0A, 0x20, 0x35), BLUE_ACCENT)
tag.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tag.text_frame.paragraphs[0].add_run()
r.text = "SOFTWARE ENGINEERING — PROJECT PROPOSAL"
set_font(r, 10, BLUE_LIGHT, True)

add_slide_number(s, 1)


# ════════════════════════════════════════════════════════════
#                SLIDE 2 — INTRODUCTION
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 2, "Introduction & Motivation")
add_slide_number(s, 2)

# Left column
tb = add_textbox(s, Inches(0.6), Inches(1.2), Inches(5.8), Inches(5.5))
tb.text_frame.word_wrap = True
first_paragraph(tb.text_frame, "🫀  The Healthcare Challenge", size=16, color=TEAL_LIGHT, bold=True)
add_bullet_list(tb.text_frame, [
    "Cardiovascular disease is the leading cause of death globally — 17.9 million deaths/year (WHO)",
    "Early detection can significantly reduce mortality rates",
    "Traditional diagnosis requires expensive imaging and expert cardiologists",
    "Machine Learning & Deep Learning can assist in risk prediction",
], size=12)

add_paragraph(tb.text_frame, "", size=6)
add_paragraph(tb.text_frame, "⚠️  The Black-Box Problem", size=16, color=AMBER, bold=True, space_before=Pt(8))
add_bullet_list(tb.text_frame, [
    "Complex AI models behave like black boxes — no explanation for decisions",
    "A prediction alone does not explain WHY a patient was classified at risk",
    "Without transparency, clinicians cannot trust or adopt AI tools",
    "Explainable AI (XAI) can make predictions understandable",
], size=12)

# Right column — flow diagrams
add_card(s, Inches(6.8), Inches(1.3), Inches(5.8), Inches(1.2),
         "❌  Current Approach", [
             "Patient Data  →  AI Model  →  Prediction Only",
             "No explanation of reasoning provided"
         ], fill=RGBColor(0x1A,0x0E,0x0E), border=RGBColor(0x5C,0x1A,0x1A),
         title_color=RED_LIGHT, body_size=12)

add_card(s, Inches(6.8), Inches(2.8), Inches(5.8), Inches(1.2),
         "✅  Proposed Approach", [
             "Patient Data  →  AI Model  →  Prediction + Explanation",
             "Transparent reasoning alongside every prediction"
         ], fill=RGBColor(0x0A,0x1A,0x18), border=RGBColor(0x14,0x5A,0x50),
         title_color=TEAL_LIGHT, body_size=12)

# Core idea box
core = add_rounded_rect(s, Inches(6.8), Inches(4.4), Inches(5.8), Inches(1.4),
                        RGBColor(0x0F,0x17,0x2A), BLUE_ACCENT)
core.text_frame.word_wrap = True
first_paragraph(core.text_frame, "Core Idea", size=14, color=BLUE_LIGHT, bold=True,
                align=PP_ALIGN.CENTER)
add_paragraph(core.text_frame, "Prediction  +  Explanation  +  Transparency",
              size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_paragraph(core.text_frame, "We propose to build a system that not only predicts,\nbut also explains its reasoning",
              size=11, color=MUTED, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
#              SLIDE 3 — PROBLEM STATEMENT
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 3, "Problem Statement")
add_slide_number(s, 3)

tb = add_textbox(s, Inches(0.6), Inches(1.15), Inches(12), Inches(0.5))
first_paragraph(tb.text_frame, "Existing heart disease prediction approaches may suffer from:", size=14, color=LIGHT_GRAY)

problems = [
    ("🔒", "Black-Box Predictions", "Models output a risk score without explaining how the decision was made", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
    ("🔍", "Limited Interpretability", "Difficult to understand which clinical features influenced a prediction", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
    ("📊", "No Visual Reasoning", "Limited visualization of model reasoning and feature contributions", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
    ("👤", "Individual Prediction Gap", "Difficult to explain why a specific patient was classified — global stats are not enough", AMBER, RGBColor(0x1A,0x15,0x05), RGBColor(0x5C,0x4A,0x0A)),
    ("⚙️", "Disconnected Systems", "Prediction and explainability tools are often developed separately", AMBER, RGBColor(0x1A,0x15,0x05), RGBColor(0x5C,0x4A,0x0A)),
]

row1_y = Inches(1.8)
card_w = Inches(3.8)
card_h = Inches(1.15)
gap = Inches(0.2)

# Row 1: 3 cards
for i, (icon, title, body, tcol, fill, brd) in enumerate(problems[:3]):
    x = Inches(0.6) + i * (card_w + gap)
    add_card(s, x, row1_y, card_w, card_h, f"{icon}  {title}", [body],
             fill=fill, border=brd, title_color=tcol, body_size=11)

# Row 2: 2 cards
for i, (icon, title, body, tcol, fill, brd) in enumerate(problems[3:]):
    x = Inches(0.6) + i * (Inches(5.9) + gap)
    add_card(s, x, row1_y + card_h + gap, Inches(5.9), card_h, f"{icon}  {title}", [body],
             fill=fill, border=brd, title_color=tcol, body_size=11)

# Research question
rq = add_rounded_rect(s, Inches(0.6), Inches(4.6), Inches(12), Inches(1.1),
                      RGBColor(0x0F,0x17,0x2A), BLUE_ACCENT)
rq.text_frame.word_wrap = True
first_paragraph(rq.text_frame, "Research Question", size=12, color=BLUE_LIGHT, bold=True)
add_paragraph(rq.text_frame,
    '"How can we design a heart disease prediction system that combines deep learning with multiple '
    'explainability techniques to provide both a prediction and understandable reasoning behind that prediction?"',
    size=13, color=LIGHT_GRAY)


# ════════════════════════════════════════════════════════════
#        SLIDE 4 — EXISTING APPROACHES & GAP
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 4, "Existing Approaches & Research Gap")
add_slide_number(s, 4)

cw = Inches(3.8)
add_card(s, Inches(0.6), Inches(1.2), cw, Inches(2.4),
         "Traditional Machine Learning", [
             "• Logistic Regression",
             "• Decision Trees & Random Forest",
             "• Support Vector Machines",
             "• K-Nearest Neighbors",
             "",
             "Interpretable but may have limited capacity for complex feature interactions"
         ], title_color=BLUE_LIGHT, body_size=11)

add_card(s, Inches(4.7), Inches(1.2), cw, Inches(2.4),
         "Deep Learning Approaches", [
             "• Convolutional Neural Networks (CNN)",
             "• Long Short-Term Memory (LSTM)",
             "• Hybrid CNN-LSTM architectures",
             "• Attention-based models",
             "",
             "Higher capacity but often function as black boxes"
         ], title_color=TEAL_LIGHT, body_size=11,
         fill=RGBColor(0x0A,0x1A,0x18), border=RGBColor(0x14,0x5A,0x50))

add_card(s, Inches(8.8), Inches(1.2), cw, Inches(2.4),
         "Explainable AI Methods", [
             "• SHAP — Shapley value attribution",
             "• LIME — Local interpretable approximations",
             "• Grad-CAM — Gradient-based visual explanations",
             "• Attention Weights — Internal focus indicators",
             "",
             "Applied individually in existing work"
         ], title_color=PURPLE_LIGHT, body_size=11,
         fill=RGBColor(0x14,0x0E,0x1E), border=RGBColor(0x4A,0x2A,0x6A))

# Gap box
add_proposed_tag(s, Inches(0.6), Inches(3.9))
add_card(s, Inches(0.6), Inches(4.35), Inches(12), Inches(2.2),
         "Proposed Research / Software Gap", [
             "This project proposes to explore a unified system combining:",
             "   • CNN-LSTM-Attention hybrid deep learning architecture",
             "   • GAF (Gramian Angular Field) representation of tabular clinical data",
             "   • Multiple XAI techniques (Attention + Grad-CAM + SHAP + LIME) in a single framework",
             "   • Integrated explainable visualization dashboard",
             "",
             "Note: The specific gap and novelty will be further validated through a comprehensive literature review."
         ], fill=RGBColor(0x1A,0x15,0x05), border=AMBER, title_color=AMBER, body_size=12)


# ════════════════════════════════════════════════════════════
#              SLIDE 5 — OBJECTIVES
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 5, "Project Objectives")
add_slide_number(s, 5)

objectives = [
    ("1", "Design a heart disease risk prediction system using deep learning that takes clinical patient parameters as input and outputs a risk assessment."),
    ("2", "Preprocess and transform clinical patient data using appropriate encoding, scaling, class balancing, and explore Gramian Angular Field (GAF) transformation for 2D representation."),
    ("3", "Explore a CNN-LSTM-Attention architecture combining spatial feature extraction (CNN), sequential relationship modeling (LSTM), and learnable feature emphasis (Attention)."),
    ("4", "Integrate multiple XAI techniques — Attention, Grad-CAM, SHAP, and LIME — to provide complementary explanations for each individual prediction."),
    ("5", "Provide understandable explanations for individual predictions by mapping AI-driven feature importance to clinically meaningful interpretations."),
    ("6", "Develop a modular software application with a web interface for prediction input, explainability visualization, and report generation."),
]

for i, (num, text) in enumerate(objectives):
    y = Inches(1.25) + i * Inches(0.95)
    # Number badge
    badge = add_rounded_rect(s, Inches(0.8), y, Inches(0.4), Inches(0.4), BLUE_ACCENT, BLUE_ACCENT)
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = badge.text_frame.paragraphs[0].add_run()
    r.text = num; set_font(r, 14, WHITE, True)

    # Text
    tb = add_textbox(s, Inches(1.4), y - Inches(0.02), Inches(10.8), Inches(0.8))
    tb.text_frame.word_wrap = True
    first_paragraph(tb.text_frame, text, size=13, color=LIGHT_GRAY)

# Footer note
tb = add_textbox(s, Inches(2), SLIDE_H - Inches(0.6), Inches(9), Inches(0.4))
first_paragraph(tb.text_frame, "All objectives are stated in planned/future tense — this is a project proposal.",
                size=10, color=MUTED, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
#           SLIDE 6 — PROPOSED SOLUTION
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 6, "Proposed Solution — Overview")
add_slide_number(s, 6)

add_proposed_tag(s, Inches(0.6), Inches(1.15))

# Left: vertical flow
flow_items = [
    ("📋  Input Patient Data", TEAL_LIGHT, RGBColor(0x0A,0x1A,0x18), RGBColor(0x14,0x5A,0x50)),
    ("⚙️  Data Preprocessing", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x25,0x63,0xEB)),
    ("🔲  GAF Transformation (1D → 2D)", PURPLE_LIGHT, RGBColor(0x14,0x0E,0x1E), RGBColor(0x4A,0x2A,0x6A)),
    ("🧠  CNN-LSTM-Attention Model", WHITE, RGBColor(0x15,0x25,0x45), RGBColor(0x3B,0x82,0xF6)),
    ("🫀  Risk Prediction", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
    ("🔍  XAI Explanation Layer", PURPLE_LIGHT, RGBColor(0x14,0x0E,0x1E), RGBColor(0x4A,0x2A,0x6A)),
    ("📊  Explainable Dashboard", AMBER, RGBColor(0x1A,0x15,0x05), RGBColor(0x5C,0x4A,0x0A)),
]

box_w = Inches(3.6)
start_y = Inches(1.65)
step = Inches(0.62)
cx = Inches(1.5)
for i, (label, tcol, fill, brd) in enumerate(flow_items):
    y = start_y + i * step
    add_flow_box(s, cx, y, box_w, label, fill, brd, tcol, 11)
    if i < len(flow_items) - 1:
        add_flow_arrow(s, cx + box_w/2 - Inches(0.2), y + Inches(0.35))

# Right: output cards
cards_data = [
    ("🎯  Risk Prediction", "Binary heart disease probability based on clinical features", BLUE_LIGHT),
    ("📊  Feature Importance", "Which clinical features contributed most to a prediction", PURPLE_LIGHT),
    ("🖼️  Visual Explanations", "Heatmaps, charts, attention maps for visual interpretation", TEAL_LIGHT),
    ("👤  Individual Interpretation", "Per-patient explanations, not just global model statistics", AMBER),
]

rhs_x = Inches(6.5)
tb = add_textbox(s, rhs_x, Inches(1.3), Inches(6), Inches(0.5))
first_paragraph(tb.text_frame, "The proposed system is intended to provide:", size=14, color=HEADING_COL, bold=True)

for i, (title, body, tcol) in enumerate(cards_data):
    y = Inches(1.85) + i * Inches(1.25)
    add_card(s, rhs_x, y, Inches(6), Inches(1.0), title, [body], title_color=tcol, body_size=12)


# ════════════════════════════════════════════════════════════
#         SLIDE 7 — SYSTEM ARCHITECTURE
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 7, "Proposed System Architecture")
add_slide_number(s, 7)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

# Main pipeline (left)
pipeline = [
    ("📁  Patient Dataset (UCI Cleveland)", TEAL_LIGHT),
    ("⚙️  Preprocessing: Impute → Scale → Encode → Balance", BLUE_LIGHT),
    ("🔲  GAF Transformation (1D → 2D Image)", PURPLE_LIGHT),
    ("📸  CNN Feature Extraction", BLUE_LIGHT),
    ("🔄  LSTM Sequential Representation", BLUE_LIGHT),
    ("🎯  Attention Mechanism", WHITE),
    ("🫀  Prediction Output", RED_LIGHT),
]

box_w = Inches(5.2)
cx = Inches(0.8)
start_y = Inches(1.65)
step = Inches(0.62)
for i, (label, tcol) in enumerate(pipeline):
    y = start_y + i * step
    fill = RGBColor(0x15,0x25,0x45) if label.startswith("🎯") else RGBColor(0x10,0x1C,0x30)
    brd = RGBColor(0x3B,0x82,0xF6) if label.startswith("🎯") else RGBColor(0x1E,0x3A,0x5F)
    add_flow_box(s, cx, y, box_w, label, fill, brd, tcol, 11)
    if i < len(pipeline) - 1:
        add_flow_arrow(s, cx + box_w/2 - Inches(0.2), y + Inches(0.35))

# Right side: XAI layer
rhs = Inches(6.8)
add_card(s, rhs, Inches(1.6), Inches(5.5), Inches(2.2),
         "Proposed XAI Layer", [
             "",
             "  Attention          Grad-CAM",
             "  SHAP                 LIME",
             "",
             "Four complementary explanation perspectives"
         ], fill=RGBColor(0x14,0x0E,0x1E), border=PURPLE, title_color=PURPLE_LIGHT, body_size=12)

add_card(s, rhs, Inches(4.1), Inches(5.5), Inches(0.9),
         "📊  Explainable Dashboard", ["Interactive visualizations, feature importance charts"],
         body_size=11)

add_card(s, rhs, Inches(5.2), Inches(5.5), Inches(0.9),
         "🗄️  MySQL Database", ["Patient records, diagnosis history, longitudinal tracking"],
         body_size=11)

add_card(s, rhs, Inches(6.3), Inches(5.5), Inches(0.7),
         "📄  Clinical PDF Reports", ["Doctor-friendly interpretation reports"],
         body_size=11)


# ════════════════════════════════════════════════════════════
#             SLIDE 8 — DATASET
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 8, "Dataset & Data Processing Plan")
add_slide_number(s, 8)

# Left: dataset info
tb = add_textbox(s, Inches(0.6), Inches(1.2), Inches(5.8), Inches(0.4))
first_paragraph(tb.text_frame, "Proposed Dataset: UCI Cleveland Heart Disease", size=15, color=HEADING_COL, bold=True)

dataset_info = [
    ("Source", "UCI Machine Learning Repository"),
    ("Samples", "~303 patient records"),
    ("Features", "13 clinical attributes + 1 target"),
    ("Target", "Binary (Normal / Heart Disease)"),
    ("Missing Values", "Some columns may require imputation"),
]
for i, (prop, val) in enumerate(dataset_info):
    y = Inches(1.7) + i * Inches(0.35)
    tb = add_textbox(s, Inches(0.8), y, Inches(5.5), Inches(0.3))
    first_paragraph(tb.text_frame, f"  {prop}:  {val}", size=12, color=LIGHT_GRAY)

# Key features
tb = add_textbox(s, Inches(0.6), Inches(3.6), Inches(5.8), Inches(0.4))
first_paragraph(tb.text_frame, "Key Clinical Features", size=14, color=HEADING_COL, bold=True)

features = [
    "age, sex, cp (chest pain), trestbps (blood pressure)",
    "chol (cholesterol), fbs (fasting blood sugar)",
    "restecg (resting ECG), thalach (max heart rate)",
    "exang (exercise angina), oldpeak (ST depression)",
    "slope (ST slope), ca (vessels), thal (thallium test)"
]
tb = add_textbox(s, Inches(0.8), Inches(4.0), Inches(5.5), Inches(2.5))
tb.text_frame.word_wrap = True
for f in features:
    add_paragraph(tb.text_frame, f"•  {f}", size=11, color=TEXT_BODY)

# Right: preprocessing plan
add_proposed_tag(s, Inches(6.8), Inches(1.2))
preprocess_steps = [
    ("1. Missing Value Handling", "Median imputation for invalid entries"),
    ("2. Feature Encoding", "One-hot encoding for categorical variables"),
    ("3. Scaling / Normalization", "Clinical-range-aware physiological scaling"),
    ("4. Outlier Analysis", "IQR-based capping to preserve dataset size"),
    ("5. Class Imbalance Handling", "SMOTE oversampling on training split only"),
    ("6. Data Splitting", "Stratified train / validation / test split"),
]
for i, (title, body) in enumerate(preprocess_steps):
    y = Inches(1.65) + i * Inches(0.88)
    add_card(s, Inches(6.8), y, Inches(5.6), Inches(0.75), title, [body], body_size=11)


# ════════════════════════════════════════════════════════════
#           SLIDE 9 — GAF TRANSFORMATION
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 9, "Gramian Angular Field (GAF) Transformation")
add_slide_number(s, 9)

# Left: Why & Steps
tb = add_textbox(s, Inches(0.6), Inches(1.2), Inches(6), Inches(0.4))
first_paragraph(tb.text_frame, "Why Consider GAF?", size=15, color=HEADING_COL, bold=True)

tb = add_textbox(s, Inches(0.6), Inches(1.6), Inches(6), Inches(2.5))
tb.text_frame.word_wrap = True
add_bullet_list(tb.text_frame, [
    "Original clinical data is tabular (1D vector)",
    "CNNs are designed for spatial pattern recognition in 2D",
    "GAF transforms a feature vector into a 2D matrix",
    "Encodes pairwise feature angular relationships",
    "The 2D texture can be explored by CNN-based extraction",
], size=12)

tb2 = add_textbox(s, Inches(0.6), Inches(3.7), Inches(6), Inches(0.4))
first_paragraph(tb2.text_frame, "Transformation Steps", size=14, color=HEADING_COL, bold=True)

steps = [
    ("Step 1", "Scale each feature to [−1, 1] range"),
    ("Step 2", "Compute angular representation: φ = arccos(x̃)"),
    ("Step 3", "Construct GASF matrix: GASF[i,j] = cos(φᵢ + φⱼ)"),
]
for i, (step_label, desc) in enumerate(steps):
    y = Inches(4.15) + i * Inches(0.65)
    add_card(s, Inches(0.6), y, Inches(6), Inches(0.55), step_label, [desc],
             title_color=BLUE_LIGHT, body_size=12)

# Right: Flow + evaluation note
rhs = Inches(7.2)
add_flow_box(s, rhs, Inches(1.5), Inches(4.5), "1D Feature Vector  (~22 values)", text_color=TEAL_LIGHT)
add_flow_arrow(s, rhs + Inches(2), Inches(1.95))
add_flow_box(s, rhs, Inches(2.3), Inches(4.5), "GAF Transformation",
             fill=RGBColor(0x15,0x25,0x45), border=RGBColor(0x3B,0x82,0xF6), text_color=WHITE)
add_flow_arrow(s, rhs + Inches(2), Inches(2.75))
add_flow_box(s, rhs, Inches(3.1), Inches(4.5), "2D Matrix  (22 × 22 × 1)",
             fill=RGBColor(0x14,0x0E,0x1E), border=PURPLE, text_color=PURPLE_LIGHT)
add_flow_arrow(s, rhs + Inches(2), Inches(3.55))
add_flow_box(s, rhs, Inches(3.9), Inches(4.5), "CNN Feature Extraction", text_color=BLUE_LIGHT)

# Conceptual grid (colored squares to represent a heatmap)
grid_x, grid_y = rhs + Inches(0.5), Inches(4.7)
cell = Inches(0.28)
colors = [
    [BLUE_ACCENT, RGBColor(0x3B,0x82,0xF6), PURPLE, TEAL, RGBColor(0x60,0xA5,0xFA), RGBColor(0x3B,0x82,0xF6), BLUE_ACCENT, RGBColor(0x22,0xD3,0xEE)],
    [RGBColor(0x3B,0x82,0xF6), PURPLE, RGBColor(0x0E,0xA5,0xE9), TEAL_LIGHT, AMBER, RGBColor(0x0E,0xA5,0xE9), PURPLE, RGBColor(0x3B,0x82,0xF6)],
    [PURPLE, RGBColor(0x0E,0xA5,0xE9), TEAL, RGBColor(0x22,0xD3,0xEE), RGBColor(0x60,0xA5,0xFA), TEAL, RGBColor(0x0E,0xA5,0xE9), PURPLE],
    [TEAL, TEAL_LIGHT, RGBColor(0x22,0xD3,0xEE), AMBER, TEAL_LIGHT, RGBColor(0x22,0xD3,0xEE), TEAL_LIGHT, TEAL],
    [RGBColor(0x60,0xA5,0xFA), AMBER, RGBColor(0x60,0xA5,0xFA), TEAL_LIGHT, BLUE_ACCENT, TEAL_LIGHT, AMBER, RGBColor(0x60,0xA5,0xFA)],
    [TEAL, TEAL_LIGHT, RGBColor(0x22,0xD3,0xEE), AMBER, TEAL_LIGHT, RGBColor(0x22,0xD3,0xEE), TEAL_LIGHT, TEAL],
    [PURPLE, RGBColor(0x0E,0xA5,0xE9), TEAL, RGBColor(0x22,0xD3,0xEE), RGBColor(0x60,0xA5,0xFA), TEAL, RGBColor(0x0E,0xA5,0xE9), PURPLE],
    [BLUE_ACCENT, RGBColor(0x3B,0x82,0xF6), PURPLE, TEAL, RGBColor(0x60,0xA5,0xFA), RGBColor(0x3B,0x82,0xF6), BLUE_ACCENT, RGBColor(0x22,0xD3,0xEE)],
]
for r, row in enumerate(colors):
    for c, clr in enumerate(row):
        sq = s.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                grid_x + c * cell, grid_y + r * cell, cell, cell)
        sq.fill.solid(); sq.fill.fore_color.rgb = clr
        sq.line.fill.background()

tb = add_textbox(s, rhs, Inches(7.0) - Inches(0.6), Inches(4.5), Inches(0.4))
first_paragraph(tb.text_frame, "Conceptual GAF Heatmap Illustration", size=10, color=MUTED, align=PP_ALIGN.CENTER)

# Evaluation note
note = add_rounded_rect(s, Inches(0.6), Inches(6.3), Inches(12), Inches(0.65),
                        RGBColor(0x0F,0x17,0x2A), BLUE_ACCENT)
note.text_frame.word_wrap = True
first_paragraph(note.text_frame,
    "The project will evaluate whether GAF-based representation improves the learning of feature interactions compared to raw tabular input.",
    size=12, color=LIGHT_GRAY)


# ════════════════════════════════════════════════════════════
#       SLIDE 10 — CNN-LSTM-ATTENTION ARCHITECTURE
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 10, "Proposed CNN-LSTM-Attention Architecture")
add_slide_number(s, 10)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

# Left: vertical flow
arch_items = [
    ("Input: GAF Representation (2D)", TEAL_LIGHT, RGBColor(0x0A,0x1A,0x18), RGBColor(0x14,0x5A,0x50)),
    ("Conv2D Layers + BatchNorm + ReLU", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x1E,0x3A,0x5F)),
    ("Spatial Feature Maps", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x1E,0x3A,0x5F)),
    ("Reshape → Sequential Format", PURPLE_LIGHT, RGBColor(0x14,0x0E,0x1E), RGBColor(0x4A,0x2A,0x6A)),
    ("LSTM / Bidirectional LSTM", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x1E,0x3A,0x5F)),
    ("Attention Mechanism", WHITE, RGBColor(0x15,0x25,0x45), RGBColor(0x3B,0x82,0xF6)),
    ("Dense + Dropout", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x1E,0x3A,0x5F)),
    ("Sigmoid → P(Heart Disease)", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
]

box_w = Inches(4)
cx = Inches(0.8)
start_y = Inches(1.55)
step = Inches(0.58)
for i, (label, tcol, fill, brd) in enumerate(arch_items):
    y = start_y + i * step
    add_flow_box(s, cx, y, box_w, label, fill, brd, tcol, 11)
    if i < len(arch_items) - 1:
        add_flow_arrow(s, cx + box_w/2 - Inches(0.2), y + Inches(0.35))

# Right: component cards
rhs = Inches(5.8)
components = [
    ("📸  CNN (Conv2D)", "Planned for spatial/local pattern extraction from the GAF image representation", BLUE_LIGHT),
    ("🔄  LSTM", "Planned to model relationships and dependencies within learned features", PURPLE_LIGHT),
    ("🎯  Attention", "Planned to emphasize important representations — outputs interpretable weights", TEAL_LIGHT),
    ("📤  Output", "Binary risk prediction — sigmoid activation producing P(Heart Disease) ∈ [0,1]", RED_LIGHT),
]
for i, (title, body, tcol) in enumerate(components):
    y = Inches(1.5) + i * Inches(1.25)
    add_card(s, rhs, y, Inches(6.5), Inches(1.0), title, [body], title_color=tcol, body_size=12)

tb = add_textbox(s, rhs, Inches(6.6), Inches(6.5), Inches(0.5))
first_paragraph(tb.text_frame, "Hyperparameters will be determined during the experimentation phase.",
                size=10, color=MUTED)


# ════════════════════════════════════════════════════════════
#            SLIDE 11 — XAI FRAMEWORK
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 11, "Multi-Technique XAI Framework")
add_slide_number(s, 11)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

xai_methods = [
    ("🎯  Attention Weights", "Will show which learned representations receive greater attention. Provides a model-internal perspective on feature importance.",
     "Intrinsic to architecture — no external computation", PURPLE_LIGHT, RGBColor(0x14,0x0E,0x1E), RGBColor(0x4A,0x2A,0x6A)),
    ("🔥  Grad-CAM", "Will provide a visual heatmap of influential regions in the CNN feature maps. Uses gradients to identify activating spatial areas.",
     "Gradient-based — requires model internals", BLUE_LIGHT, RGBColor(0x0F,0x17,0x2A), RGBColor(0x25,0x63,0xEB)),
    ("🧮  SHAP", "Will estimate each feature's marginal contribution toward an individual prediction using Shapley values from game theory.",
     "Model-agnostic — mathematically grounded", TEAL_LIGHT, RGBColor(0x0A,0x1A,0x18), RGBColor(0x14,0x5A,0x50)),
    ("🔍  LIME", "Will provide a local interpretable approximation by fitting a simple model on perturbed neighborhood samples around a prediction.",
     "Model-agnostic — perturbation-based", AMBER, RGBColor(0x1A,0x15,0x05), RGBColor(0x5C,0x4A,0x0A)),
]

cw = Inches(5.8)
for i, (title, body, note, tcol, fill, brd) in enumerate(xai_methods):
    col = i % 2
    row = i // 2
    x = Inches(0.6) + col * (cw + Inches(0.3))
    y = Inches(1.6) + row * Inches(2.15)
    add_card(s, x, y, cw, Inches(1.9), title, [body, "", note], fill=fill, border=brd,
             title_color=tcol, body_size=12)

# Bottom: consensus note
note = add_rounded_rect(s, Inches(0.6), Inches(5.95), Inches(12), Inches(0.75),
                        RGBColor(0x0F,0x17,0x2A), BLUE_ACCENT)
note.text_frame.word_wrap = True
first_paragraph(note.text_frame, "One Prediction → Four Explanation Perspectives", size=13, color=BLUE_LIGHT, bold=True,
                align=PP_ALIGN.CENTER)
add_paragraph(note.text_frame, "The project will investigate whether these methods provide complementary or consistent explanations.",
              size=12, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
#          SLIDE 12 — SOFTWARE ARCHITECTURE
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 12, "Software Architecture")
add_slide_number(s, 12)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

# Left: stack
stack_items = [
    ("👨‍⚕️  User / Doctor", TEAL_LIGHT),
    ("🌐  Web Interface (HTML / JS / Plotly)", BLUE_LIGHT),
    ("⚡  Flask Application (API)", WHITE),
    ("🔮  Prediction Service", BLUE_LIGHT),
    ("🧠  CNN-LSTM-Attention Model", PURPLE_LIGHT),
    ("🔍  XAI Engine", PURPLE_LIGHT),
    ("📊  Visualization  +  🗄️ MySQL DB", AMBER),
]
cx = Inches(0.8)
box_w = Inches(4.6)
for i, (label, tcol) in enumerate(stack_items):
    y = Inches(1.65) + i * Inches(0.62)
    highlight = label.startswith("⚡")
    fill = RGBColor(0x15,0x25,0x45) if highlight else RGBColor(0x10,0x1C,0x30)
    brd = RGBColor(0x3B,0x82,0xF6) if highlight else RGBColor(0x1E,0x3A,0x5F)
    add_flow_box(s, cx, y, box_w, label, fill, brd, tcol, 11)
    if i < len(stack_items) - 1:
        add_flow_arrow(s, cx + box_w/2 - Inches(0.2), y + Inches(0.35))

# Right: modules table + tech
rhs = Inches(6.2)
tb = add_textbox(s, rhs, Inches(1.4), Inches(6), Inches(0.4))
first_paragraph(tb.text_frame, "Planned Software Modules", size=14, color=HEADING_COL, bold=True)

modules = [
    ("User Input Module", "Patient data entry form"),
    ("Preprocessing Module", "Cleaning, scaling, GAF"),
    ("Prediction Module", "CNN-LSTM-Attention inference"),
    ("XAI Module", "Attention, Grad-CAM, SHAP, LIME"),
    ("Visualization Module", "Plotly charts & heatmaps"),
    ("Database Module", "Patient records (MySQL)"),
    ("Report Module", "PDF clinical report generation"),
]
for i, (mod, desc) in enumerate(modules):
    y = Inches(1.85) + i * Inches(0.55)
    tb = add_textbox(s, rhs, y, Inches(6), Inches(0.45))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    r1 = p.add_run(); r1.text = f"  {mod}  "; set_font(r1, 11, WHITE, True)
    r2 = p.add_run(); r2.text = f"—  {desc}"; set_font(r2, 11, MUTED)

tb = add_textbox(s, rhs, Inches(5.8), Inches(6), Inches(0.4))
first_paragraph(tb.text_frame, "Planned Technologies", size=14, color=HEADING_COL, bold=True)

techs = "Python  •  TensorFlow/Keras  •  Flask  •  MySQL  •  SHAP  •  LIME  •  Plotly  •  Pandas  •  Scikit-learn  •  NumPy"
tb = add_textbox(s, rhs, Inches(6.2), Inches(6), Inches(0.6))
tb.text_frame.word_wrap = True
first_paragraph(tb.text_frame, techs, size=11, color=BLUE_LIGHT)


# ════════════════════════════════════════════════════════════
#          SLIDE 13 — DEVELOPMENT METHODOLOGY
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 13, "Development Methodology")
add_slide_number(s, 13)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

phases = [
    ("P1", "Requirement Analysis", "Stakeholder needs, feature specification, feasibility", BLUE_ACCENT),
    ("P2", "Dataset Collection & Preprocessing", "Data cleaning, encoding, scaling, GAF, SMOTE", BLUE_ACCENT),
    ("P3", "Baseline Model Development", "CNN-only baseline for initial benchmarking", PURPLE),
    ("P4", "CNN-LSTM-Attention Development", "Full proposed architecture implementation", PURPLE),
    ("P5", "XAI Integration", "Attention, Grad-CAM, SHAP, LIME implementation", TEAL),
    ("P6", "Web Application Development", "Flask backend, dashboard UI, MySQL integration", TEAL),
    ("P7", "Testing & Evaluation", "Metrics, ablation study, XAI consistency analysis", AMBER),
    ("P8", "Documentation & Final Deployment", "Report writing, PPT, demo preparation", AMBER),
]

for i, (pnum, title, desc, color) in enumerate(phases):
    col = i % 2
    row = i // 2
    x = Inches(0.6) + col * Inches(6.3)
    y = Inches(1.55) + row * Inches(1.3)

    # Phase badge
    badge = add_rounded_rect(s, x, y, Inches(0.5), Inches(0.5), color, color)
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = badge.text_frame.paragraphs[0].add_run()
    r.text = pnum; set_font(r, 11, WHITE, True)

    # Card
    add_card(s, x + Inches(0.65), y, Inches(5.4), Inches(1.05), title, [desc],
             body_size=11, title_size=13)


# ════════════════════════════════════════════════════════════
#         SLIDE 14 — TESTING & EVALUATION PLAN
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 14, "Testing & Evaluation Plan")
add_slide_number(s, 14)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

# Top-left: Metrics
add_card(s, Inches(0.6), Inches(1.6), Inches(5.8), Inches(2.4),
         "📊  Model Evaluation Metrics", [
             "Performance will be evaluated using:",
             "  • Accuracy — overall prediction correctness",
             "  • Precision — false positive control",
             "  • Recall — false negative control (critical in healthcare)",
             "  • F1-Score — precision-recall balance",
             "  • ROC-AUC — discrimination ability",
             "  • Confusion Matrix — detailed error breakdown",
         ], title_color=BLUE_LIGHT, body_size=11)

# Top-right: Validation
add_card(s, Inches(6.7), Inches(1.6), Inches(5.8), Inches(2.4),
         "✅  Validation Strategy", [
             "  • Hold-out testing (70/10/20 stratified split)",
             "  • Stratified K-Fold cross-validation for robustness",
             "  • Ablation study: CNN vs CNN-LSTM vs CNN-LSTM-Attention",
             "  • SMOTE applied inside each fold (no data leakage)",
             "",
             "The model will be compared against baseline variants."
         ], title_color=TEAL_LIGHT, body_size=11,
         fill=RGBColor(0x0A,0x1A,0x18), border=RGBColor(0x14,0x5A,0x50))

# Bottom-left: XAI
add_card(s, Inches(0.6), Inches(4.3), Inches(5.8), Inches(2.4),
         "🔍  XAI Evaluation", [
             "  • Feature ranking comparison across XAI methods",
             "  • Top-K feature overlap (Jaccard Similarity)",
             "  • Rank correlation (Spearman) between methods",
             "  • Consistency check — do methods agree?",
             "  • Comparison with established medical literature",
         ], title_color=PURPLE_LIGHT, body_size=11,
         fill=RGBColor(0x14,0x0E,0x1E), border=RGBColor(0x4A,0x2A,0x6A))

# Bottom-right: Software
add_card(s, Inches(6.7), Inches(4.3), Inches(5.8), Inches(2.4),
         "🧪  Software Testing", [
             "  • Unit testing — preprocessing, GAF, model functions",
             "  • Integration testing — end-to-end API pipeline",
             "  • Input validation — handling invalid/missing inputs",
             "  • Edge-case testing — out-of-distribution values",
             "  • UI testing — form validation, result display",
         ], title_color=AMBER, body_size=11,
         fill=RGBColor(0x1A,0x15,0x05), border=RGBColor(0x5C,0x4A,0x0A))


# ════════════════════════════════════════════════════════════
#           SLIDE 15 — EXPECTED OUTCOMES
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_header(s, 15, "Expected Outcomes")
add_slide_number(s, 15)
add_proposed_tag(s, Inches(0.6), Inches(1.15))

deliverables = [
    ("🧠", "Prediction Model", "A trained CNN-LSTM-Attention model for heart disease risk prediction", BLUE_LIGHT),
    ("🔲", "GAF Pipeline", "Preprocessing pipeline converting tabular data to 2D representations", PURPLE_LIGHT),
    ("🔍", "XAI Module", "Multi-technique explainability: Attention, Grad-CAM, SHAP & LIME", TEAL_LIGHT),
    ("📊", "Dashboard", "Interactive web visualization for predictions and explanations", AMBER),
    ("🌐", "Flask Application", "Modular Flask-based software system with MySQL backend", BLUE_LIGHT),
    ("📋", "Evaluation Report", "Model comparison, ablation study, XAI consistency analysis", RED_LIGHT),
]

for i, (icon, title, body, tcol) in enumerate(deliverables):
    col = i % 3
    row = i // 3
    x = Inches(0.6) + col * Inches(4.1)
    y = Inches(1.6) + row * Inches(1.85)
    add_card(s, x, y, Inches(3.8), Inches(1.6), f"{icon}  {title}", [body],
             title_color=tcol, body_size=12)

# Expected output flow
flow_box = add_rounded_rect(s, Inches(0.6), Inches(5.5), Inches(12), Inches(1.2),
                            RGBColor(0x0F,0x17,0x2A), BLUE_ACCENT)
flow_box.text_frame.word_wrap = True
first_paragraph(flow_box.text_frame, "Expected System Output", size=13, color=BLUE_LIGHT, bold=True,
                align=PP_ALIGN.CENTER)
add_paragraph(flow_box.text_frame,
    "Patient Input  →  Risk Prediction  →  Feature Importance  →  Visual Explanation",
    size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
add_paragraph(flow_box.text_frame,
    "No numerical performance claims — results will be evaluated during development.",
    size=10, color=MUTED, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
#        SLIDE 16 — FUTURE SCOPE & LIMITATIONS
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_MID)
add_slide_header(s, 16, "Future Scope & Current Limitations")
add_slide_number(s, 16)

# Left: future scope
tb = add_textbox(s, Inches(0.6), Inches(1.2), Inches(5.8), Inches(0.4))
first_paragraph(tb.text_frame, "🚀  Future Scope", size=15, color=HEADING_COL, bold=True)

future_items = [
    "📂  Larger & diverse datasets (Framingham, MIMIC-III)",
    "🌍  External dataset validation for generalization",
    "🔬  Additional XAI methods (Integrated Gradients, Counterfactuals)",
    "🤖  Advanced architectures (Vision Transformers)",
    "📱  Mobile deployment for clinical use",
    "📈  Longitudinal patient risk progression tracking",
    "🔒  Privacy-preserving federated learning",
]
for i, item in enumerate(future_items):
    y = Inches(1.7) + i * Inches(0.7)
    add_card(s, Inches(0.6), y, Inches(5.8), Inches(0.58), item, [],
             body_size=11, title_size=12, title_color=LIGHT_GRAY)

# Right: limitations
tb = add_textbox(s, Inches(6.8), Inches(1.2), Inches(5.8), Inches(0.4))
first_paragraph(tb.text_frame, "⚠️  Current Scope & Limitations", size=15, color=AMBER, bold=True)

limitations = [
    ("Dataset Size", "Initial development uses a public dataset (~303 records). Size may limit generalization."),
    ("XAI Interpretation", "Feature importance does not always imply direct causation. Explanations require careful clinical interpretation."),
    ("Not a Medical Device", "The system is a research/educational prototype, NOT an autonomous diagnosis tool. Clinical validation required for real deployment."),
    ("Single Dataset", "Conclusions will be specific to Cleveland dataset. Multi-dataset validation is planned as future work."),
]
for i, (title, body) in enumerate(limitations):
    y = Inches(1.7) + i * Inches(1.25)
    fill = RGBColor(0x1A,0x0E,0x0E) if title == "Not a Medical Device" else RGBColor(0x1A,0x15,0x05)
    brd = RGBColor(0x5C,0x1A,0x1A) if title == "Not a Medical Device" else RGBColor(0x5C,0x4A,0x0A)
    tcol = RED_LIGHT if title == "Not a Medical Device" else AMBER
    add_card(s, Inches(6.8), y, Inches(5.8), Inches(1.05), title, [body],
             fill=fill, border=brd, title_color=tcol, body_size=11)


# ════════════════════════════════════════════════════════════
#           SLIDE 17 — CONCLUSION & THANK YOU
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s, BG_DARK)
add_slide_number(s, 17)

# Decorative line
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.8), Inches(1.2), Inches(1.7), Pt(3))
line.fill.solid(); line.fill.fore_color.rgb = BLUE_ACCENT; line.line.fill.background()

# Section label
tb = add_textbox(s, Inches(2), Inches(1.5), Inches(9.3), Inches(0.4))
first_paragraph(tb.text_frame, "CONCLUSION", size=12, color=HEADING_COL, bold=True, align=PP_ALIGN.CENTER)

# Summary text
tb = add_textbox(s, Inches(1.5), Inches(2.0), Inches(10.3), Inches(1.2))
tb.text_frame.word_wrap = True
first_paragraph(tb.text_frame,
    "The proposed project aims to develop an explainable heart disease prediction system combining:",
    size=16, color=LIGHT_GRAY, align=PP_ALIGN.CENTER)

# Component badges row
components = [
    ("Deep Learning", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x25,0x63,0xEB)),
    ("GAF", PURPLE_LIGHT, RGBColor(0x14,0x0E,0x1E), RGBColor(0x4A,0x2A,0x6A)),
    ("Attention", TEAL_LIGHT, RGBColor(0x0A,0x1A,0x18), RGBColor(0x14,0x5A,0x50)),
    ("Grad-CAM", BLUE_LIGHT, RGBColor(0x10,0x1C,0x30), RGBColor(0x25,0x63,0xEB)),
    ("SHAP", AMBER, RGBColor(0x1A,0x15,0x05), RGBColor(0x5C,0x4A,0x0A)),
    ("LIME", RED_LIGHT, RGBColor(0x1A,0x0E,0x0E), RGBColor(0x5C,0x1A,0x1A)),
]

total_w = len(components) * Inches(1.7) + (len(components) - 1) * Inches(0.25)
start_x = (SLIDE_W - total_w) / 2
for i, (label, tcol, fill, brd) in enumerate(components):
    x = start_x + i * (Inches(1.7) + Inches(0.25))
    add_flow_box(s, x, Inches(3.3), Inches(1.7), label, fill, brd, tcol, 11)

# Plus signs between
for i in range(len(components) - 1):
    x = start_x + Inches(1.7) + i * (Inches(1.7) + Inches(0.25))
    tb = add_textbox(s, x, Inches(3.3), Inches(0.25), Inches(0.38))
    first_paragraph(tb.text_frame, "+", size=14, color=BLUE_ACCENT, bold=True, align=PP_ALIGN.CENTER)

# Goal text
tb = add_textbox(s, Inches(1.5), Inches(4.1), Inches(10.3), Inches(0.8))
tb.text_frame.word_wrap = True
first_paragraph(tb.text_frame,
    "The goal is to investigate whether combining prediction with multiple explanation techniques\ncan make AI outputs more transparent and understandable.",
    size=13, color=MUTED, align=PP_ALIGN.CENTER)

# ECG line
ecg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(5.0), Inches(4.3), Pt(2))
ecg.fill.solid(); ecg.fill.fore_color.rgb = RGBColor(0x1E,0x3A,0x5F); ecg.line.fill.background()

# Thank You
tb = add_textbox(s, Inches(2), Inches(5.3), Inches(9.3), Inches(0.9))
first_paragraph(tb.text_frame, "Thank You", size=40, color=WHITE, bold=True,
                align=PP_ALIGN.CENTER, name='Calibri')

tb = add_textbox(s, Inches(2), Inches(6.2), Inches(9.3), Inches(0.5))
first_paragraph(tb.text_frame, "Questions & Suggestions", size=16, color=BLUE_LIGHT,
                align=PP_ALIGN.CENTER)


# ═══════════════════ SAVE ═══════════════════
output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                           "Heart_Disease_Project_Proposal.pptx")
prs.save(output_path)
print(f"✅ Presentation saved to: {output_path}")
print(f"   Total slides: {len(prs.slides)}")
