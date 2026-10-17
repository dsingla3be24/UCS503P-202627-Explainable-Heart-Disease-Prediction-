import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Custom color palette (CardioAI Theme)
    DARK_NAVY = RGBColor(10, 15, 30)      # Background
    TEXT_WHITE = RGBColor(248, 250, 252)  # Primary Text
    CORAL_PINK = RGBColor(255, 99, 132)  # Secondary Accent (Highlight)
    LIGHT_BLUE = RGBColor(56, 189, 248)  # Secondary Accent (Info)
    MUTED_GRAY = RGBColor(148, 163, 184) # Muted Text
    CARD_BG = RGBColor(24, 30, 48)       # Background for info panels/cards

    # Helper function to apply dark background
    def apply_dark_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = DARK_NAVY

    # Helper function to add slide title
    def add_slide_header(slide, title_text):
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = LIGHT_BLUE
        return title_box

    # Helper function to add card/shape container
    def add_card(slide, left, top, width, height):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = LIGHT_BLUE
        shape.line.width = Pt(1)
        return shape

    blank_layout = prs.slide_layouts[6] # Blank slide layout

    # ==================== SLIDE 1: TITLE SLIDE ====================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide1)
    
    # Big Title
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "CardioAI"
    p.font.name = "Segoe UI"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = CORAL_PINK
    p.alignment = PP_ALIGN.LEFT
    
    p2 = tf.add_paragraph()
    p2.text = "An Explainable Clinical Decision Support System (CDSS) for Heart Disease Diagnosis"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(22)
    p2.font.color.rgb = TEXT_WHITE
    p2.font.bold = True
    
    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(11.33), Inches(1.8))
    tf_sub = sub_box.text_frame
    p3 = tf_sub.paragraphs[0]
    p3.text = "Transforming 1D Tabular Clinical Data into 2D GAF Images with CNN-LSTM-Attention & Multi-Model XAI Interpretation"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(16)
    p3.font.italic = True
    p3.font.color.rgb = LIGHT_BLUE
    
    p4 = tf_sub.add_paragraph()
    p4.text = "\nB.Tech Project Presentation  |  Presenter: Student Representative"
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(14)
    p4.font.color.rgb = MUTED_GRAY

    # ==================== SLIDE 2: CLINICAL CONTEXT ====================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide2)
    add_slide_header(slide2, "Clinical Context & Problem Statement")
    
    # Left Card (The Problem)
    add_card(slide2, Inches(0.75), Inches(1.8), Inches(5.6), Inches(4.8))
    text_box_left = slide2.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_l = text_box_left.text_frame
    tf_l.word_wrap = True
    p_l_t = tf_l.paragraphs[0]
    p_l_t.text = "The Heart Disease Challenge"
    p_l_t.font.name = "Segoe UI"
    p_l_t.font.size = Pt(20)
    p_l_t.font.bold = True
    p_l_t.font.color.rgb = CORAL_PINK
    
    points_l = [
        "Cardiovascular Diseases (CVDs) are the #1 killer globally, responsible for 32% of all deaths (17.9M annually).",
        "Clinical checkups generate valuable tabular variables (blood pressure, cholesterol, resting ECG parameters).",
        "Subtle correlations and interactions between these features are complex and often missed during manual review.",
        "Clinical diagnosis is highly time-consuming, subjective, and prone to errors under high workload."
    ]
    for pt in points_l:
        p = tf_l.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right Card (Limitations of Existing AI Solutions)
    add_card(slide2, Inches(6.98), Inches(1.8), Inches(5.6), Inches(4.8))
    text_box_right = slide2.shapes.add_textbox(Inches(7.13), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_r = text_box_right.text_frame
    tf_r.word_wrap = True
    p_r_t = tf_r.paragraphs[0]
    p_r_t.text = "Limitations of Existing AI Solutions"
    p_r_t.font.name = "Segoe UI"
    p_r_t.font.size = Pt(20)
    p_r_t.font.bold = True
    p_r_t.font.color.rgb = LIGHT_BLUE
    
    points_r = [
        "Black Box Problem: Standard Machine Learning models (like XGBoost/SVM) predict risk but do not explain 'Why', leading to zero trust from medical staff.",
        "Lack of Production Context: Theoretical research models run inside Python notebooks without real EHR databases, user forms, or dashboard interfaces.",
        "No Quality Checks: Models fail to check fairness (accuracy disparities between genders and age subgroups)."
    ]
    for pt in points_r:
        p = tf_r.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # ==================== SLIDE 3: SYSTEM METHODOLOGY ====================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide3)
    add_slide_header(slide3, "System Architecture: Tabular to 2D GASF Image")
    
    # Main Concept Box
    add_card(slide3, Inches(0.75), Inches(1.8), Inches(11.83), Inches(4.8))
    text_box_m = slide3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(4.4))
    tf_m = text_box_m.text_frame
    tf_m.word_wrap = True
    
    p = tf_m.paragraphs[0]
    p.text = "Why Convert Tabular Data to a 2D GASF Image?"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CORAL_PINK
    p.space_after = Pt(14)
    
    p2 = tf_m.add_paragraph()
    p2.text = "1. Enabling 2D CNNs on Tabular Data: 2D Convolutional layers (Conv2D) extract spatial relationships. To use them, 1D metrics (28 features post one-hot encoding) are transformed into a 28x28 grid."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(8)
    
    p3 = tf_m.add_paragraph()
    p3.text = "2. Polar Coordinate Mapping: Features scaled to [-1, 1] are converted to angles using phi = arccos(x). This preserves absolute relationships without distortion."
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_WHITE
    p3.space_after = Pt(8)
    
    p4 = tf_m.add_paragraph()
    p4.text = "3. Gramian Angular Summation Field (GASF) Equation:"
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(15)
    p4.font.bold = True
    p4.font.color.rgb = LIGHT_BLUE
    
    p_math = tf_m.add_paragraph()
    p_math.text = "    GASF[i, j] = cos(phi_i + phi_j) = x_i * x_j - sqrt(1 - x_i^2) * sqrt(1 - x_j^2)"
    p_math.font.name = "Courier New"
    p_math.font.size = Pt(16)
    p_math.font.bold = True
    p_math.font.color.rgb = CORAL_PINK
    p_math.space_after = Pt(8)
    
    p5 = tf_m.add_paragraph()
    p5.text = "4. Automatic Feature Interactions: Every cell (i, j) in the 2D image represents the interaction between feature i and feature j. The 2D CNN automatically learns these combinations during training."
    p5.font.name = "Segoe UI"
    p5.font.size = Pt(15)
    p5.font.color.rgb = TEXT_WHITE

    # ==================== SLIDE 4: NETWORK PIPELINE ====================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide4)
    add_slide_header(slide4, "Deep learning Network: DCNN + Bi-LSTM + Attention")
    
    # 3-step pipeline columns
    widths = Inches(3.68)
    gap = Inches(0.4)
    start_left = Inches(0.75)
    
    steps = [
        {
            "title": "1. 2D CNN (Spatial)",
            "color": CORAL_PINK,
            "text": "• Inputs the 28x28x1 GASF image.\n• Uses Conv2D filters (128 & 256 size) with LeakyReLU activations.\n• Extracts spatial feature patterns representing the immediate relationships between adjacent variables (like blood pressure and ECG oldpeak)."
        },
        {
            "title": "2. Bi-LSTM (Temporal)",
            "color": LIGHT_BLUE,
            "text": "• Reshapes the pooled CNN feature maps back into a sequence (28 steps).\n• Bidirectional LSTM (256 units forward + 256 units backward) extracts sequential relationships across features.\n• Captures long-range, complex global dependencies."
        },
        {
            "title": "3. Tabular Attention",
            "color": TEXT_WHITE,
            "text": "• Projects Bi-LSTM outputs through a projection weight matrix (W) and bias (b).\n• Uses Softmax normalization to calculate custom relevance attention scores (alphas) for each feature.\n• Outputs a weighted context vector representing high-impact clinical drivers."
        }
    ]
    
    for i, step in enumerate(steps):
        left_pos = start_left + i * (widths + gap)
        add_card(slide4, left_pos, Inches(1.8), widths, Inches(4.8))
        tb = slide4.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.0), widths - Inches(0.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = step["title"]
        p.font.name = "Segoe UI"
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = step["color"]
        p.space_after = Pt(14)
        
        p2 = tf.add_paragraph()
        p2.text = step["text"]
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_WHITE
        p2.space_after = Pt(10)

    # ==================== SLIDE 5: XAI SUITE ====================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide5)
    add_slide_header(slide5, "Multi-Model Explainable AI (XAI) Suite")
    
    add_card(slide5, Inches(0.75), Inches(1.8), Inches(11.83), Inches(4.8))
    tb_xai = slide5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(4.4))
    tf_xai = tb_xai.text_frame
    tf_xai.word_wrap = True
    
    p = tf_xai.paragraphs[0]
    p.text = "Three Pillars of Diagnostic Trust"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CORAL_PINK
    p.space_after = Pt(14)
    
    p1 = tf_xai.add_paragraph()
    p1.text = "🔥 Grad-CAM (CNN Layer Activation): Highlights spatial regions within the 2D correlation matrix that triggered the CNN filters. Shows feature-interaction regions."
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(15)
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)
    
    p2 = tf_xai.add_paragraph()
    p2.text = "🎯 Attention Weights (Bi-LSTM Step Relevance): Calculates the step-wise weights (alphas) applied during sequence integration. Identifies which specific attributes standardly drove the prediction sequence."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(8)
    
    p3 = tf_xai.add_paragraph()
    p3.text = "⚖️ SHAP Values (Shapley Additive exPlanations): Uses Game Theory (KernelExplainer wrapper) to measure how much each attribute shifted the final probability compared to the base value."
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_WHITE
    p3.space_after = Pt(14)
    
    p4 = tf_xai.add_paragraph()
    p4.text = "Aggregation Engine: Dummy variables created by one-hot encoding are dynamically aggregated back to the 13 original clinical categories for easier interpretation by doctors."
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(15)
    p4.font.bold = True
    p4.font.color.rgb = LIGHT_BLUE

    # ==================== SLIDE 6: WEB APPLICATION ====================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide6)
    add_slide_header(slide6, "Clinical CDSS Console: Production Features")
    
    # Left Card
    add_card(slide6, Inches(0.75), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_web_l = slide6.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_w_l = tb_web_l.text_frame
    tf_w_l.word_wrap = True
    p_w_l = tf_w_l.paragraphs[0]
    p_w_l.text = "Interactive Clinician Console"
    p_w_l.font.name = "Segoe UI"
    p_w_l.font.size = Pt(20)
    p_w_l.font.bold = True
    p_w_l.font.color.rgb = CORAL_PINK
    
    points_wl = [
        "Flask-Powered Medical Dashboard: Simple UI for entering patient variables (blood pressure, cholesterol, stress-test metrics).",
        "Plotly Interactive Attributions: Renders side-by-side bar charts dynamically comparing SHAP, Grad-CAM, and Attention weights.",
        "Live GASF Heatmap rendering: Showcases the 2D angular representation of the patient on-screen.",
        "Patient search console: Integrated with electronic records to query history via patient name or ID."
    ]
    for pt in points_wl:
        p = tf_w_l.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right Card
    add_card(slide6, Inches(6.98), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_web_r = slide6.shapes.add_textbox(Inches(7.13), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_w_r = tb_web_r.text_frame
    tf_w_r.word_wrap = True
    p_w_r = tf_w_r.paragraphs[0]
    p_w_r.text = "Algorithmic Fairness (Anti-Bias)"
    p_w_r.font.name = "Segoe UI"
    p_w_r.font.size = Pt(20)
    p_w_r.font.bold = True
    p_w_r.font.color.rgb = LIGHT_BLUE
    
    points_wr = [
        "Clinically Unbiased AI: Medical diagnostics models must be validated across subgroups (Genders, Ages).",
        "Dynamic Subgroup Metrics: Calculates accuracy, precision, and specificity metrics on the fly for similar demographic groups based on patient age and sex.",
        "Allows doctors to verify that the model is performing equitably for both male and female patients, checking validation consistency."
    ]
    for pt in points_wr:
        p = tf_w_r.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # ==================== SLIDE 7: PDF REPORTS ====================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide7)
    add_slide_header(slide7, "Automated Doctor-Friendly PDF Reports")
    
    add_card(slide7, Inches(0.75), Inches(1.8), Inches(11.83), Inches(4.8))
    tb_pdf = slide7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(4.4))
    tf_pdf = tb_pdf.text_frame
    tf_pdf.word_wrap = True
    
    p = tf_pdf.paragraphs[0]
    p.text = "Dynamic Report Generation Engine (`src/pdf_generator.py`)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CORAL_PINK
    p.space_after = Pt(14)
    
    p1 = tf_pdf.add_paragraph()
    p1.text = "• Clinical Data Summary: Neatly formats raw physiological inputs, binarized targets, and risk scores in a single page document."
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(15)
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)
    
    p2 = tf_pdf.add_paragraph()
    p2.text = "• Embedded Heatmaps & Charts: Converts matplotlib outputs into temporary files and inserts the 2D GAF map alongside XAI bar charts in the PDF."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(8)
    
    p3 = tf_pdf.add_paragraph()
    p3.text = "• Decision Recommendations: Parses patient age and maximum heart rate to generate guidelines-based cardiac suggestions."
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_WHITE
    p3.space_after = Pt(8)
    
    p4 = tf_pdf.add_paragraph()
    p4.text = "• Signature & Clinician blocks: Provides standard sign-off segments for medical record integration, allowing printable export directly to hospital records."
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(15)
    p4.font.color.rgb = TEXT_WHITE

    # ==================== SLIDE 8: RESEARCH GAP ====================
    slide8 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide8)
    add_slide_header(slide8, "Research Gaps in Paper & CardioAI Uniqueness")
    
    # Left Card
    add_card(slide8, Inches(0.75), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_gap_l = slide8.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_g_l = tb_gap_l.text_frame
    tf_g_l.word_wrap = True
    p_g_l = tf_g_l.paragraphs[0]
    p_g_l.text = "Gaps In Gaurav Kumar (2025) Paper"
    p_g_l.font.name = "Segoe UI"
    p_g_l.font.size = Pt(20)
    p_g_l.font.bold = True
    p_g_l.font.color.rgb = CORAL_PINK
    
    points_gl = [
        "Mismatch Title: Promotes continuous 'ECG signals analysis' but uses simple static tabular UCI Cleveland data.",
        "Static Sequence Fallacy: Claims Bi-LSTM models 'temporal health dependencies over time' on a dataset where each patient has exactly one snapshot.",
        "Zero Real-world Deployment: Stalls at Python script metrics; no EHR connections, visual dashboards, or printable reports."
    ]
    for pt in points_gl:
        p = tf_g_l.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right Card
    add_card(slide8, Inches(6.98), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_gap_r = slide8.shapes.add_textbox(Inches(7.13), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_g_r = tb_gap_r.text_frame
    tf_g_r.word_wrap = True
    p_g_r = tf_g_r.paragraphs[0]
    p_g_r.text = "CardioAI Unique Improvements"
    p_g_r.font.name = "Segoe UI"
    p_g_r.font.size = Pt(20)
    p_g_r.font.bold = True
    p_g_r.font.color.rgb = LIGHT_BLUE
    
    points_gr = [
        "Ablation Studies Implementations: Complete script for CNN, CNN-LSTM, and CNN-LSTM-Attention comparative verification.",
        "Fail-Safe Engine: Emulation pipeline to fallback to scikit-learn MLP Classifier if native TensorFlow DLL fails.",
        "Production-Grade application: Dynamic UI console, dynamic Plotly visualizations, MySQL database storage, and PDF generation."
    ]
    for pt in points_gr:
        p = tf_g_r.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # ==================== SLIDE 9: DATABASE INTEGRATION ====================
    slide9 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide9)
    add_slide_header(slide9, "EHR & Database Architecture (`src/models.py`)")
    
    add_card(slide9, Inches(0.75), Inches(1.8), Inches(11.83), Inches(4.8))
    tb_db = slide9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.33), Inches(4.4))
    tf_db = tb_db.text_frame
    tf_db.word_wrap = True
    
    p = tf_db.paragraphs[0]
    p.text = "MySQL Database Integration & History Tracking"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = CORAL_PINK
    p.space_after = Pt(14)
    
    p1 = tf_db.add_paragraph()
    p1.text = "• Patient Management (`patients` table): Auto-registers new patients, generating clean hospital IDs. Links repeat patient names with single database IDs."
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(15)
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)
    
    p2 = tf_db.add_paragraph()
    p2.text = "• Clinical Records (`clinical_records` table): Stores age, sex, chest pain parameters, ST-segment measurements, and records date/time."
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_after = Pt(8)
    
    p3 = tf_db.add_paragraph()
    p3.text = "• AI Diagnoses (`ai_diagnoses` table): Tracks diagnosis probability, prediction risk labels, top attention driving features, and local paths of saved PDF reports."
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = TEXT_WHITE
    p3.space_after = Pt(8)
    
    p4 = tf_db.add_paragraph()
    p4.text = "• Longitudinal Patient Tracking: Automatically pulls up the patient's previous measurements during a search, calculating age increases dynamically over time."
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(15)
    p4.font.bold = True
    p4.font.color.rgb = LIGHT_BLUE

    # ==================== SLIDE 10: SUMMARY ====================
    slide10 = prs.slides.add_slide(blank_layout)
    apply_dark_background(slide10)
    add_slide_header(slide10, "Summary & Future Scope")
    
    # Left Card
    add_card(slide10, Inches(0.75), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_s_l = slide10.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_s_l = tb_s_l.text_frame
    tf_s_l.word_wrap = True
    p_s_l = tf_s_l.paragraphs[0]
    p_s_l.text = "Key Deliverables"
    p_s_l.font.name = "Segoe UI"
    p_s_l.font.size = Pt(20)
    p_s_l.font.bold = True
    p_s_l.font.color.rgb = CORAL_PINK
    
    points_sl = [
        "Engineered a production-ready Clinical Decision Support System (CDSS) for heart disease.",
        "Successfully implemented dynamic 2D GASF conversion and multi-XAI visualization (Grad-CAM, Attention, SHAP).",
        "Achieved a functional fail-safe deployment using scikit-learn MLP Classifier fallback.",
        "Created an automated doctor-friendly clinical PDF generator and database EHR lookup."
    ]
    for pt in points_sl:
        p = tf_s_l.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right Card
    add_card(slide10, Inches(6.98), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_s_r = slide10.shapes.add_textbox(Inches(7.13), Inches(1.9), Inches(5.3), Inches(4.6))
    tf_s_r = tb_s_r.text_frame
    tf_s_r.word_wrap = True
    p_s_r = tf_s_r.paragraphs[0]
    p_s_r.text = "Future Work"
    p_s_r.font.name = "Segoe UI"
    p_s_r.font.size = Pt(20)
    p_s_r.font.bold = True
    p_s_r.font.color.rgb = LIGHT_BLUE
    
    points_sr = [
        "Multi-Center Database integration: Extend validation to Hungary, Switzerland, and Statlog clinical datasets.",
        "IoT ECG Smartwatch integration: Capture raw ECG signal waves dynamically via smartwatches to replace static clinical data.",
        "Federated Learning: Enable collaborative training on hospital networks without exposing patient records."
    ]
    for pt in points_sr:
        p = tf_s_r.add_paragraph()
        p.text = "• " + pt
        p.font.name = "Segoe UI"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Save presentation
    output_path = os.path.abspath("CardioAI_Project_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")
    return output_path

if __name__ == '__main__':
    create_presentation()
